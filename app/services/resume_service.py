from __future__ import annotations

import hashlib
import io
import re
import uuid
import warnings
import zipfile
from dataclasses import dataclass
from pathlib import Path, PurePath
from typing import Any, BinaryIO

from bs4 import BeautifulSoup
from docx import Document
from flask import current_app
from odf import teletype
from odf.opendocument import load as load_odt
from PIL import Image
from pypdf import PdfReader
from pypdf.errors import PdfReadError
from striprtf.striprtf import rtf_to_text
from werkzeug.datastructures import FileStorage
from werkzeug.utils import secure_filename

from ..extensions import db
from ..models import Resume, User
from ..utils import normalize_phone


class ResumeValidationError(ValueError):
    """A resume failed a user-correctable safety or format check."""


@dataclass(frozen=True)
class ValidatedResume:
    original_name: str
    safe_name: str
    extension: str
    mime_type: str
    data: bytes
    sha256: str


@dataclass(frozen=True)
class ResumeParseResult:
    text: str
    profile: dict[str, Any]
    status: str
    message: str


MAX_ZIP_ENTRIES = 2_000
MAX_UNCOMPRESSED_BYTES = 60 * 1024 * 1024
MAX_COMPRESSION_RATIO = 150
MAX_EXTRACTED_CHARACTERS = 250_000

TEXT_EXTENSIONS = {"txt", "md"}
HTML_EXTENSIONS = {"html", "htm"}
IMAGE_EXTENSIONS = {"jpg", "jpeg", "png", "webp"}

MIME_BY_EXTENSION = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "doc": "application/msword",
    "odt": "application/vnd.oasis.opendocument.text",
    "rtf": "application/rtf",
    "txt": "text/plain",
    "md": "text/markdown",
    "html": "text/html",
    "htm": "text/html",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "png": "image/png",
    "webp": "image/webp",
}

SKILL_ALIASES = {
    "python": "Python",
    "java": "Java",
    "javascript": "JavaScript",
    "js": "JavaScript",
    "typescript": "TypeScript",
    "react": "React",
    "reactjs": "React",
    "angular": "Angular",
    "vue": "Vue.js",
    "vuejs": "Vue.js",
    "node": "Node.js",
    "nodejs": "Node.js",
    "flask": "Flask",
    "django": "Django",
    "fastapi": "FastAPI",
    "sql": "SQL",
    "mysql": "MySQL",
    "postgresql": "PostgreSQL",
    "postgres": "PostgreSQL",
    "mongodb": "MongoDB",
    "html": "HTML",
    "css": "CSS",
    "git": "Git",
    "github": "GitHub",
    "docker": "Docker",
    "kubernetes": "Kubernetes",
    "aws": "AWS",
    "azure": "Azure",
    "gcp": "Google Cloud",
    "linux": "Linux",
    "excel": "Excel",
    "power bi": "Power BI",
    "powerbi": "Power BI",
    "tableau": "Tableau",
    "pandas": "Pandas",
    "numpy": "NumPy",
    "scikit-learn": "scikit-learn",
    "sklearn": "scikit-learn",
    "tensorflow": "TensorFlow",
    "pytorch": "PyTorch",
    "machine learning": "Machine Learning",
    "data analysis": "Data Analysis",
    "data structures": "Data Structures",
    "algorithms": "Algorithms",
    "rest api": "REST APIs",
    "restful api": "REST APIs",
    "figma": "Figma",
    "ui/ux": "UI/UX",
    "communication": "Communication",
    "problem solving": "Problem Solving",
}

EDUCATION_TERMS = (
    "bachelor",
    "b.tech",
    "btech",
    "b.e.",
    "master",
    "m.tech",
    "mtech",
    "mba",
    "bca",
    "mca",
    "degree",
    "university",
    "college",
    "institute",
)

EXPERIENCE_SECTION_HEADINGS = (
    "experience",
    "work experience",
    "professional experience",
    "internship",
    "internships",
    "projects",
    "project experience",
)

SECTION_STOP_HEADINGS = (
    "education",
    "skills",
    "technical skills",
    "certifications",
    "achievements",
    "awards",
    "summary",
    "objective",
    "languages",
    "interests",
)


def _read_bounded(stream: BinaryIO, limit: int) -> bytes:
    data = stream.read(limit + 1)
    if len(data) > limit:
        raise ResumeValidationError(f"Resume must be smaller than {limit // (1024 * 1024)} MB.")
    if not data:
        raise ResumeValidationError("The selected resume is empty.")
    return data


def _validate_zip(data: bytes, extension: str) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                raise ResumeValidationError("The document contains too many embedded files.")
            total_size = 0
            for entry in entries:
                total_size += entry.file_size
                if total_size > MAX_UNCOMPRESSED_BYTES:
                    raise ResumeValidationError("The document expands beyond the safe size limit.")
                compressed = max(1, entry.compress_size)
                if (
                    entry.file_size > 1_000_000
                    and entry.file_size / compressed > MAX_COMPRESSION_RATIO
                ):
                    raise ResumeValidationError("The document has an unsafe compression ratio.")
                path = PurePath(entry.filename.replace("\\", "/"))
                if path.is_absolute() or ".." in path.parts:
                    raise ResumeValidationError("The document contains unsafe internal paths.")
            names = {entry.filename for entry in entries}
            if extension == "docx" and not {
                "[Content_Types].xml",
                "word/document.xml",
            }.issubset(names):
                raise ResumeValidationError("This file is not a valid DOCX document.")
            if extension == "odt":
                if "mimetype" not in names:
                    raise ResumeValidationError("This file is not a valid ODT document.")
                media_type = archive.read("mimetype").decode("ascii", errors="ignore")
                if media_type.strip() != MIME_BY_EXTENSION["odt"]:
                    raise ResumeValidationError("This file is not a valid ODT document.")
    except zipfile.BadZipFile as exc:
        raise ResumeValidationError(
            f"This file is not a valid {extension.upper()} document."
        ) from exc


def _validate_image(data: bytes, extension: str) -> None:
    expected = {"jpg": "JPEG", "jpeg": "JPEG", "png": "PNG", "webp": "WEBP"}[extension]
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(io.BytesIO(data)) as image:
                if image.format != expected:
                    raise ResumeValidationError(
                        "The image content does not match its file extension."
                    )
                image.verify()
    except ResumeValidationError:
        raise
    except (Image.DecompressionBombError, Image.DecompressionBombWarning) as exc:
        raise ResumeValidationError("The resume image dimensions are too large.") from exc
    except (OSError, ValueError) as exc:
        raise ResumeValidationError("This is not a valid resume image.") from exc


def validate_resume_upload(
    upload: FileStorage,
    *,
    allowed_extensions: set[str] | None = None,
    max_bytes: int | None = None,
) -> ValidatedResume:
    filename = Path(upload.filename or "").name
    if not filename or "." not in filename:
        raise ResumeValidationError("Choose a resume file with a supported extension.")
    extension = filename.rsplit(".", 1)[1].lower()
    allowed = allowed_extensions or set(current_app.config["ALLOWED_RESUME_EXTENSIONS"])
    if extension not in allowed:
        raise ResumeValidationError("That resume format is not supported.")

    size_limit = int(max_bytes or current_app.config.get("MAX_CONTENT_LENGTH", 10 * 1024 * 1024))
    try:
        data = _read_bounded(upload.stream, size_limit)
    finally:
        try:
            upload.stream.seek(0)
        except (OSError, AttributeError):
            pass

    if extension == "pdf" and not data.lstrip().startswith(b"%PDF-"):
        raise ResumeValidationError("This file is not a valid PDF document.")
    if extension in {"docx", "odt"}:
        _validate_zip(data, extension)
    elif extension == "doc" and not data.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        raise ResumeValidationError("This file is not a valid legacy Word document.")
    elif extension == "rtf" and not data.lstrip().lower().startswith(b"{\\rtf"):
        raise ResumeValidationError("This file is not a valid RTF document.")
    elif extension in IMAGE_EXTENSIONS:
        _validate_image(data, extension)
    elif extension in TEXT_EXTENSIONS | HTML_EXTENSIONS:
        if b"\x00" in data[:4096]:
            raise ResumeValidationError("This text resume contains invalid binary content.")
        try:
            data.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise ResumeValidationError("Text resumes must use UTF-8 encoding.") from exc

    cleaned_name = secure_filename(filename) or f"resume.{extension}"
    return ValidatedResume(
        original_name=filename[:255],
        safe_name=cleaned_name[:255],
        extension=extension,
        mime_type=MIME_BY_EXTENSION.get(extension, "application/octet-stream"),
        data=data,
        sha256=hashlib.sha256(data).hexdigest(),
    )


def _normalize_text(text: str) -> str:
    text = text.replace("\x00", " ").replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    cleaned = "\n".join(line for line in lines if line)
    return cleaned[:MAX_EXTRACTED_CHARACTERS]


def _extract_pdf(data: bytes) -> tuple[str, str]:
    try:
        reader = PdfReader(io.BytesIO(data), strict=False)
        if reader.is_encrypted:
            try:
                unlocked = reader.decrypt("")
            except Exception as exc:
                raise ResumeValidationError("Password-protected PDFs cannot be parsed.") from exc
            if not unlocked:
                raise ResumeValidationError("Password-protected PDFs cannot be parsed.")
        max_pages = int(current_app.config.get("RESUME_MAX_PAGES", 20))
        if len(reader.pages) > max_pages:
            raise ResumeValidationError(f"Resume PDFs may contain at most {max_pages} pages.")
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    except ResumeValidationError:
        raise
    except (PdfReadError, OSError, ValueError) as exc:
        raise ResumeValidationError("The PDF is damaged or cannot be read.") from exc
    cleaned = _normalize_text(text)
    status = "parsed" if len(cleaned) >= 60 else "needs_ocr"
    return cleaned, status


def _extract_docx(data: bytes) -> str:
    try:
        document = Document(io.BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return _normalize_text("\n".join(parts))
    except (ValueError, KeyError, zipfile.BadZipFile) as exc:
        raise ResumeValidationError("The DOCX document could not be parsed.") from exc


def _extract_odt(data: bytes) -> str:
    try:
        document = load_odt(io.BytesIO(data))
        return _normalize_text(teletype.extractText(document.text))
    except Exception as exc:
        raise ResumeValidationError("The ODT document could not be parsed.") from exc


def _decode_utf8(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise ResumeValidationError("Text resumes must use UTF-8 encoding.") from exc


def extract_resume_text(data: bytes, extension: str) -> tuple[str, str, str]:
    """Return ``text, status, message`` without inventing content for OCR/DOC files."""

    extension = extension.lower().lstrip(".")
    if extension == "pdf":
        text, status = _extract_pdf(data)
        message = (
            "Resume parsed successfully."
            if status == "parsed"
            else "This PDF appears scanned. Upload DOCX/TXT or enter your details manually."
        )
        return text, status, message
    if extension == "docx":
        text = _extract_docx(data)
    elif extension == "odt":
        text = _extract_odt(data)
    elif extension == "rtf":
        try:
            text = _normalize_text(rtf_to_text(data.decode("latin-1")))
        except (UnicodeError, ValueError) as exc:
            raise ResumeValidationError("The RTF document could not be parsed.") from exc
    elif extension in TEXT_EXTENSIONS:
        text = _normalize_text(_decode_utf8(data))
    elif extension in HTML_EXTENSIONS:
        soup = BeautifulSoup(_decode_utf8(data), "html.parser")
        for unsafe in soup(["script", "style", "template", "noscript"]):
            unsafe.decompose()
        text = _normalize_text(soup.get_text("\n"))
    elif extension == "doc":
        return "", "unsupported", (
            "Legacy .doc parsing is not available safely. Save the file as DOCX or PDF and retry."
        )
    elif extension in IMAGE_EXTENSIONS:
        return "", "needs_ocr", (
            "Image OCR is not enabled. Upload a text-based PDF/DOCX or enter details manually."
        )
    else:
        raise ResumeValidationError("That resume format is not supported by the parser.")

    if len(text) < 40:
        return (
            text,
            "needs_review",
            "Very little text was found; please complete your profile manually.",
        )
    return text, "parsed", "Resume parsed successfully. Review the details before saving."


def extract_skills_from_text(text: str) -> list[str]:
    lowered = re.sub(r"[^a-z0-9+#./-]+", " ", text.lower())
    found: set[str] = set()
    for alias, canonical in SKILL_ALIASES.items():
        pattern = rf"(?<![a-z0-9]){re.escape(alias)}(?![a-z0-9])"
        if re.search(pattern, lowered):
            found.add(canonical)
    return sorted(found, key=str.casefold)


def _section_heading(line: str) -> str:
    return re.sub(r"[^a-z ]+", " ", line.casefold()).strip()


def _extract_section(
    lines: list[str],
    headings: tuple[str, ...],
    *,
    stop_headings: tuple[str, ...],
    limit: int = 16,
) -> list[str]:
    """Extract a bounded resume section without treating the whole file as experience."""

    collected: list[str] = []
    collecting = False
    all_headings = set(headings) | set(stop_headings) | set(EXPERIENCE_SECTION_HEADINGS)
    for raw_line in lines:
        line = raw_line.strip()
        heading = _section_heading(line)
        looks_like_heading = len(line) <= 80 and heading in all_headings
        if heading in headings and len(line) <= 80:
            collecting = True
            continue
        if collecting and looks_like_heading:
            break
        if collecting:
            collected.append(line)
            if len(collected) >= limit:
                break
    return collected


def parse_profile(text: str) -> dict[str, Any]:
    cleaned = _normalize_text(text)
    lines = [line for line in cleaned.splitlines() if line]

    email_match = re.search(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", cleaned, re.I)
    phone_value = ""
    for candidate in re.findall(r"(?:\+?\d[\d ()-]{8,}\d)", cleaned):
        try:
            phone_value = normalize_phone(candidate)
            break
        except ValueError:
            continue

    name = ""
    headline = ""
    for line in lines[:8]:
        plain = line.strip(" |-•")
        lower = plain.casefold()
        if not plain or "@" in plain or re.search(r"\d{4,}", plain):
            continue
        if lower in {"resume", "curriculum vitae", "cv", "profile", "summary"}:
            continue
        if not name and 1 < len(plain.split()) <= 5 and len(plain) <= 80:
            name = plain
            continue
        if not headline and len(plain) <= 160:
            headline = plain
            break

    education_lines = [
        line
        for line in lines
        if any(term in line.casefold() for term in EDUCATION_TERMS)
    ]
    years = [int(year) for year in re.findall(r"\b(?:19|20)\d{2}\b", cleaned)]
    graduation_year = max((year for year in years if 1980 <= year <= 2100), default=None)

    experience_values = [
        float(value)
        for value in re.findall(
            r"\b(\d{1,2}(?:\.\d+)?)\s*(?:\+\s*)?(?:years?|yrs?)\b",
            cleaned,
            re.I,
        )
    ]
    experience_years = max(experience_values, default=0.0)
    if re.search(r"\b(fresher|entry[- ]level|no experience)\b", cleaned, re.I):
        experience_years = 0.0

    experience_lines: list[str] = []
    for section_heading in EXPERIENCE_SECTION_HEADINGS:
        for line in _extract_section(
            lines,
            (section_heading,),
            stop_headings=SECTION_STOP_HEADINGS,
            limit=12,
        ):
            if line not in experience_lines:
                experience_lines.append(line)
            if len(experience_lines) >= 20:
                break
        if len(experience_lines) >= 20:
            break

    links = []
    for url in re.findall(r"https?://[^\s<>()]+", cleaned, re.I):
        cleaned_url = url.rstrip(".,;:)")
        if cleaned_url not in links:
            links.append(cleaned_url)

    return {
        "full_name": name,
        "email": email_match.group(0).lower() if email_match else "",
        "phone": phone_value,
        "headline": headline,
        "education": "\n".join(education_lines[:8]),
        "experience": "\n".join(experience_lines[:20]),
        "skills": extract_skills_from_text(cleaned),
        "experience_years": experience_years,
        "graduation_year": graduation_year,
        "links": links[:10],
    }


def parse_resume_bytes(data: bytes, extension: str) -> ResumeParseResult:
    text, status, message = extract_resume_text(data, extension)
    profile = parse_profile(text) if text else {
        "full_name": "",
        "email": "",
        "phone": "",
        "headline": "",
        "education": "",
        "experience": "",
        "skills": [],
        "experience_years": 0.0,
        "graduation_year": None,
        "links": [],
    }
    return ResumeParseResult(text=text, profile=profile, status=status, message=message)


def save_and_parse_resume(
    user: User,
    upload: FileStorage,
    *,
    commit: bool = True,
) -> Resume:
    from .file_cleanup import (
        is_discarded_private_file,
        process_private_file,
        queue_private_file,
    )

    validated = validate_resume_upload(upload)
    parsed = parse_resume_bytes(validated.data, validated.extension)
    storage_mode = current_app.config.get("RESUME_STORAGE_MODE", "local")
    destination: Path | None = None
    generated_name = f"{uuid.uuid4().hex}.{validated.extension}"
    if storage_mode == "local":
        upload_root = Path(current_app.config["UPLOAD_FOLDER"]).resolve()
        user_folder = (upload_root / str(user.id)).resolve()
        if upload_root not in user_folder.parents:
            raise ResumeValidationError("Unable to create a safe resume storage path.")
        user_folder.mkdir(parents=True, exist_ok=True)
        destination = (user_folder / generated_name).resolve()
        if user_folder not in destination.parents:
            raise ResumeValidationError("Unable to create a safe resume storage path.")
        destination.write_bytes(validated.data)
        stored_name = f"{user.id}/{generated_name}"
    else:
        # Serverless filesystems are ephemeral. Parse in memory and retain only
        # structured, user-reviewable profile data and non-content metadata.
        stored_name = f"discarded/{user.id}/{generated_name}"

    record = Resume(
        user_id=user.id,
        original_name=validated.original_name,
        stored_name=stored_name,
        mime_type=validated.mime_type,
        extension=validated.extension,
        size_bytes=len(validated.data),
        sha256=validated.sha256,
        extracted_text=parsed.text if storage_mode == "local" else "",
        parsed_profile=parsed.profile,
        parse_status=parsed.status,
        parse_message=parsed.message,
        is_current=True,
    )
    previous_records = db.session.scalars(
        db.select(Resume).where(Resume.user_id == user.id)
    ).all()
    cleanup_ids: list[int] = []
    try:
        cleanup_records = [
            queue_private_file(previous.stored_name)
            for previous in previous_records
            if not is_discarded_private_file(previous.stored_name)
        ]
        for previous in previous_records:
            db.session.delete(previous)
        db.session.add(record)
        db.session.flush()
        cleanup_ids = [cleanup.id for cleanup in cleanup_records]
        if commit:
            db.session.commit()
        else:
            db.session.flush()
    except Exception:
        db.session.rollback()
        if destination is not None:
            destination.unlink(missing_ok=True)
        raise
    if commit:
        for cleanup_id in cleanup_ids:
            process_private_file(cleanup_id)
    return record


# Compatibility aliases used by thin route modules.
validate_and_parse_resume = save_and_parse_resume
process_resume_upload = save_and_parse_resume
